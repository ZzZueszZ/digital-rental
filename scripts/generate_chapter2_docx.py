from __future__ import annotations

import copy
import os
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DOCX = Path(r"D:\PERSONAL\hoc-ki-2-nam-4\KLTN\KLTN-22162027-22162022_TruongAiNga-NguyenThanhLoc.docx")
OUT_DIR = ROOT / "docs" / "generated" / "chapter-2"
DIAGRAM_DIR = OUT_DIR / "diagrams"
OUT_DOCX = OUT_DIR / "KLTN-chuong-2-phan-tich-thiet-ke.docx"


def ensure_dirs() -> None:
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)


def write_diagrams() -> list[tuple[str, str, str]]:
    """Return (slug, title, mmd_path)."""
    diagrams: list[tuple[str, str, str]] = [
        (
            "01-architecture",
            "Sơ đồ kiến trúc tổng quan hệ thống",
            r"""
flowchart LR
    Browser["Web Browser"] --> FE["Next.js Frontend"]
    FE --> API["Spring Boot REST API"]
    API --> PG["PostgreSQL Database"]
    API --> Redis["Redis Token Blacklist / Cache"]
    API --> Mail["SMTP Mail Service"]
    API --> VNPay["VNPay Payment Gateway"]
    API --> Uploads["Uploaded Identity / Product Files"]
""",
        ),
        (
            "02-usecase-overview",
            "Sơ đồ Use Case tổng quát",
            r"""
flowchart LR
    Guest["Guest"]
    Customer["Customer"]
    StaffAdmin["Staff / Admin"]
    SuperAdmin["Super Admin"]
    VNPay["VNPay"]
    Mail["Mail Service"]

    subgraph System["Camera Rental Web Application"]
        UC1(("Browse catalog"))
        UC2(("Register / Login"))
        UC3(("Manage profile and address"))
        UC4(("Cart and purchase checkout"))
        UC5(("VNPay payment"))
        UC6(("Rental lifecycle"))
        UC7(("eKYC submission and review"))
        UC8(("Catalog and inventory management"))
        UC9(("Order and voucher management"))
        UC10(("Review and support management"))
        UC11(("User / Role / Permission management"))
        UC12(("Dashboard and audit logs"))
    end

    Guest --> UC1
    Guest --> UC2
    Guest --> UC10
    Customer --> UC2
    Customer --> UC3
    Customer --> UC4
    Customer --> UC5
    Customer --> UC6
    Customer --> UC7
    StaffAdmin --> UC7
    StaffAdmin --> UC8
    StaffAdmin --> UC9
    StaffAdmin --> UC10
    StaffAdmin --> UC12
    SuperAdmin --> UC11
    SuperAdmin --> UC12
    VNPay --> UC5
    Mail --> UC2
""",
        ),
        (
            "03-activity-auth",
            "Activity Diagram đăng ký, kích hoạt và đăng nhập",
            r"""
flowchart TD
    Start([Start]) --> Register["Customer submits registration form"]
    Register --> Validate{"Valid data?"}
    Validate -- No --> ShowError["Show validation error"]
    ShowError --> Register
    Validate -- Yes --> CreateUser["Backend creates inactive account"]
    CreateUser --> SendMail["Mail service sends activation link"]
    SendMail --> Activate["Customer opens activation link"]
    Activate --> ActiveOk{"Token valid?"}
    ActiveOk -- No --> Expired["Show expired / invalid link"]
    ActiveOk -- Yes --> MarkActive["Backend activates account"]
    MarkActive --> Login["Customer logs in"]
    Login --> AuthOk{"Credentials valid?"}
    AuthOk -- No --> LoginError["Show login error"]
    AuthOk -- Yes --> IssueToken["Issue access and refresh token"]
    IssueToken --> End([End])
""",
        ),
        (
            "04-activity-checkout",
            "Activity Diagram mua hàng và thanh toán",
            r"""
flowchart TD
    Start([Start]) --> Browse["Browse product"]
    Browse --> AddCart["Add item to cart"]
    AddCart --> Cart["Review cart"]
    Cart --> Voucher["Apply voucher optional"]
    Voucher --> Address{"Shipping address available?"}
    Address -- No --> AddAddress["Add shipping address"]
    AddAddress --> Checkout
    Address -- Yes --> Checkout["Create order"]
    Checkout --> Stock{"Stock available?"}
    Stock -- No --> Fail["Return checkout error"]
    Stock -- Yes --> Payment["Create VNPay payment URL"]
    Payment --> Pay["Customer pays on VNPay"]
    Pay --> Callback["VNPay returns result"]
    Callback --> Verify{"Signature and result valid?"}
    Verify -- No --> PaymentFail["Mark payment failed"]
    Verify -- Yes --> PaymentOk["Update order payment status"]
    PaymentOk --> Track["Customer tracks order"]
    Track --> End([End])
""",
        ),
        (
            "05-activity-rental",
            "Activity Diagram vòng đời đơn thuê",
            r"""
flowchart TD
    Start([Start]) --> Check["Check rental availability"]
    Check --> Available{"Available?"}
    Available -- No --> Stop["Choose another date/device"]
    Available -- Yes --> Checkout["Create rental order"]
    Checkout --> Contract["Customer signs rental contract"]
    Contract --> Fee["Customer pays rental fee"]
    Fee --> Prepare["Staff prepares device"]
    Prepare --> Report["Staff creates handover report"]
    Report --> Deposit["Staff collects deposit"]
    Deposit --> Handover["Staff hands over device"]
    Handover --> Return["Customer returns device"]
    Return --> ReturnReport["Staff creates return report"]
    ReturnReport --> Complete["Staff completes rental"]
    Complete --> End([End])
""",
        ),
        (
            "06-sequence-auth",
            "Sequence Diagram đăng ký, kích hoạt và đăng nhập",
            r"""
sequenceDiagram
    actor Customer
    participant FE as Frontend
    participant API as Backend API
    participant DB as PostgreSQL
    participant Mail as Mail Service
    Customer->>FE: Submit registration
    FE->>API: POST /auth/register
    API->>DB: Create inactive user and activation token
    API->>Mail: Send activation email
    Mail-->>Customer: Activation link
    Customer->>FE: Open activation link
    FE->>API: GET /auth/activate
    API->>DB: Activate account
    Customer->>FE: Submit login
    FE->>API: POST /auth/login
    API->>DB: Verify user and roles
    API-->>FE: Access token, refresh token, user profile
""",
        ),
        (
            "07-sequence-vnpay",
            "Sequence Diagram thanh toán VNPay",
            r"""
sequenceDiagram
    actor Customer
    participant FE as Frontend
    participant API as Backend API
    participant VNPay
    participant DB as PostgreSQL
    Customer->>FE: Click pay
    FE->>API: POST /payments/vnpay/create
    API->>API: Build signed payment request
    API-->>FE: Payment URL
    FE-->>Customer: Redirect to VNPay
    Customer->>VNPay: Complete payment
    VNPay->>API: GET /payments/vnpay/return
    API->>API: Verify secure hash and result
    API->>DB: Update order/payment log
    API-->>FE: Redirect payment result page
""",
        ),
        (
            "08-sequence-ekyc",
            "Sequence Diagram eKYC khách hàng và quản trị duyệt",
            r"""
sequenceDiagram
    actor Customer
    actor Admin
    participant FE as Frontend
    participant API as Backend API
    participant DB as PostgreSQL
    participant FileStore as Upload Storage
    Customer->>FE: Start eKYC
    FE->>API: POST /ekyc/initiate
    API->>DB: Create verification session
    Customer->>FE: Upload front, back, selfie
    FE->>API: POST /ekyc/upload-front, upload-back, upload-selfie
    API->>FileStore: Store artifacts
    API->>DB: Link artifacts to session
    Customer->>FE: Submit eKYC
    FE->>API: POST /ekyc/submit
    API->>DB: Mark session pending
    Admin->>FE: Open pending eKYC list
    FE->>API: GET /admin/ekyc/pending
    API-->>FE: Pending sessions
    Admin->>FE: Approve or reject
    FE->>API: PUT /admin/ekyc/{id}/approve or reject
    API->>DB: Update verification status
""",
        ),
        (
            "09-sequence-rental",
            "Sequence Diagram vòng đời thuê thiết bị",
            r"""
sequenceDiagram
    actor Customer
    actor Staff
    participant FE as Frontend
    participant API as Backend API
    participant DB as PostgreSQL
    participant VNPay
    Customer->>FE: Check rental availability
    FE->>API: GET /rentals/products/{id}/availability
    API->>DB: Query devices and rental schedule
    API-->>FE: Availability result
    Customer->>FE: Checkout rental
    FE->>API: POST /rentals/checkout
    API->>DB: Create rental order and contract
    Customer->>FE: Sign contract
    FE->>API: POST /rentals/{id}/contract/sign
    API->>DB: Update contract status
    FE->>API: POST /payments/vnpay/rental-fee/create
    API-->>FE: Rental fee payment URL
    Customer->>VNPay: Pay rental fee
    VNPay->>API: GET /payments/vnpay/rental-fee/return
    API->>DB: Update rental payment
    Staff->>FE: Prepare and handover device
    FE->>API: POST prepare, handover-report, collect-deposit, handover
    API->>DB: Update rental lifecycle
    Staff->>FE: Process return
    FE->>API: POST return-report and complete
    API->>DB: Complete rental order
""",
        ),
        (
            "10-sequence-admin-catalog-inventory",
            "Sequence Diagram quản lý catalog và tồn kho",
            r"""
sequenceDiagram
    actor Admin
    participant FE as Admin UI
    participant API as Backend API
    participant DB as PostgreSQL
    Admin->>FE: Create or update product
    FE->>API: POST /products or PUT /products/{id}/info
    API->>DB: Save product, category, images
    Admin->>FE: Update price
    FE->>API: PUT /products/{id}/price
    API->>DB: Save price history
    Admin->>FE: Adjust stock
    FE->>API: PUT /inventory/products/{productId}/stock
    API->>DB: Update product stock
    API->>DB: Write inventory audit log
    Admin->>FE: View low stock/dashboard
    FE->>API: GET /dashboard/low-stock
    API-->>FE: Low stock data
""",
        ),
        (
            "11-sequence-rbac",
            "Sequence Diagram quản lý role và permission",
            r"""
sequenceDiagram
    actor SuperAdmin
    participant FE as Super Admin UI
    participant API as Backend API
    participant DB as PostgreSQL
    SuperAdmin->>FE: Create permission
    FE->>API: POST /permissions
    API->>DB: Save permission
    SuperAdmin->>FE: Create role
    FE->>API: POST /roles
    API->>DB: Save role
    SuperAdmin->>FE: Assign permission to role
    FE->>API: POST /roles/{id}/permissions
    API->>DB: Update role-permission mapping
    SuperAdmin->>FE: Update user role/status
    FE->>API: PUT /users/{id} or PATCH /users/{id}/status
    API->>DB: Update user
""",
        ),
        (
            "12-erd-logical",
            "ERD mức logic của các nhóm dữ liệu chính",
            r"""
erDiagram
    USER ||--|| USER_PROFILE : has
    USER ||--o{ SHIPPING_ADDRESS : owns
    USER ||--o{ ORDER : places
    USER ||--o{ RENTAL_ORDER : rents
    USER ||--o{ REVIEW : writes
    USER ||--o{ VERIFICATION_SESSION : submits
    ROLE }o--o{ PERMISSION : grants
    USER }o--o{ ROLE : assigned
    CATEGORY ||--o{ PRODUCT : contains
    PRODUCT ||--o{ PRODUCT_IMAGE : has
    PRODUCT ||--o{ PRODUCT_PRICE_HISTORY : tracks
    PRODUCT ||--o{ INVENTORY_AUDIT_LOG : changes
    ORDER ||--o{ ORDER_ITEM : includes
    PRODUCT ||--o{ ORDER_ITEM : sold
    RENTAL_ORDER ||--o{ RENTAL_ORDER_ITEM : includes
    PRODUCT ||--o{ DEVICE : has
    DEVICE ||--o{ RENTAL_ORDER_ITEM : assigned
    RENTAL_ORDER ||--|| RENTAL_CONTRACT : has
    RENTAL_ORDER ||--o{ RENTAL_PAYMENT : paid_by
    RENTAL_ORDER ||--o{ RENTAL_HANDOVER_REPORT : handover
    RENTAL_ORDER ||--o{ RENTAL_RETURN_REPORT : return
    VOUCHER ||--o{ ORDER : applies
    VERIFICATION_SESSION ||--o{ VERIFICATION_ARTIFACT : contains
    VERIFICATION_SESSION ||--o{ VERIFICATION_RESULT : produces
    SUPPORT_TICKET }o--|| USER : optional_owner
    AUDIT_LOG }o--|| USER : actor
""",
        ),
        (
            "13-sitemap",
            "Sơ đồ điều hướng chính của hệ thống",
            r"""
flowchart TD
    Home["/"] --> ProductDetail["/products/{id}"]
    Home --> Login["/auth/login"]
    Login --> Register["/auth/register"]
    Login --> Forgot["/auth/forgot-password"]
    Home --> Checkout["/checkout"]
    Checkout --> VNPayReturn["/checkout/vnpay-return"]
    Home --> Profile["/profile"]
    Profile --> ProfileInfo["/profile/info"]
    Profile --> ProfileAddress["/profile/address"]
    Profile --> ProfileCart["/profile/cart"]
    Profile --> ProfileOrders["/profile/orders"]
    Profile --> ProfileEkyc["/profile/ekyc"]
    Home --> RentalReturn["/rentals/payment-return"]
    Login --> Admin["/admin | /staff | /super-admin"]
    Admin --> Dashboard["dashboard"]
    Admin --> Users["users"]
    Admin --> Products["products"]
    Admin --> Categories["categories"]
    Admin --> Orders["orders"]
    Admin --> Rentals["rentals"]
    Admin --> Vouchers["vouchers"]
    Admin --> Reviews["reviews"]
    Admin --> Support["support"]
    Admin --> EkycAdmin["ekyc"]
    Admin --> Audit["audit-logs"]
    Admin --> RBAC["roles / permissions"]
""",
        ),
    ]
    for slug, _, content in diagrams:
        path = DIAGRAM_DIR / f"{slug}.mmd"
        path.write_text(content.strip() + "\n", encoding="utf-8")
    return [(slug, title, str(DIAGRAM_DIR / f"{slug}.mmd")) for slug, title, _ in diagrams]


def export_diagrams(diagrams: list[tuple[str, str, str]]) -> list[tuple[str, str, Path]]:
    exported: list[tuple[str, str, Path]] = []
    npx = shutil.which("npx.cmd") or shutil.which("npx") or shutil.which("npx.ps1")
    if not npx:
        raise RuntimeError("Cannot find npx. Install Node.js/npm or add npx to PATH.")
    for slug, title, mmd_path in diagrams:
        out = DIAGRAM_DIR / f"{slug}.png"
        if out.exists() and out.stat().st_size > 0:
            exported.append((slug, title, out))
            continue
        cmd = [
            npx,
            "--yes",
            "@mermaid-js/mermaid-cli",
            "-i",
            mmd_path,
            "-o",
            str(out),
            "-b",
            "white",
        ]
        subprocess.run(cmd, cwd=ROOT, check=True)
        exported.append((slug, title, out))
    return exported


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(15 if level <= 1 else 13)
    if level <= 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_para(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0.3)
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(13)
        r2 = p.add_run(text[len(bold_prefix) :])
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(13)
    else:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style=None)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        r = p.add_run("• " + item)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(11)


def add_figure(doc: Document, title: str, path: Path, number: int) -> int:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.2))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(f"Hình 2.{number}. {title}")
    r.italic = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    return number + 1


def build_chapter2(doc: Document, diagrams: list[tuple[str, str, Path]]) -> None:
    fig = 1
    diagram = {slug: (title, path) for slug, title, path in diagrams}

    add_heading(doc, "Chương 2: PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG", 1)

    add_heading(doc, "2.1. PHÂN TÍCH NGHIỆP VỤ HỆ THỐNG", 2)
    add_para(
        doc,
        "Hệ thống được xây dựng cho bài toán thuê và mua thiết bị máy ảnh trực tuyến. Người dùng có thể tìm kiếm thiết bị, xem thông tin sản phẩm, đặt mua, đặt thuê, thanh toán trực tuyến và theo dõi trạng thái đơn. Ở phía quản trị, hệ thống hỗ trợ quản lý sản phẩm, danh mục, tồn kho, đơn mua, đơn thuê, voucher, người dùng, eKYC, hỗ trợ khách hàng và nhật ký thao tác.",
    )
    add_para(
        doc,
        "Đặc thù của bài toán là thiết bị cho thuê có giá trị cao, vì vậy hệ thống cần kiểm soát danh tính khách hàng, lịch sử thuê, hợp đồng, tiền cọc và trạng thái bàn giao/trả thiết bị. Bên cạnh đó, dữ liệu xử lý trong hệ thống có độ nhạy cảm cao như thông tin cá nhân, địa chỉ, ảnh giấy tờ định danh, khuôn mặt, hợp đồng thuê và lịch sử thanh toán.",
    )
    add_table(
        doc,
        ["Nhóm nghiệp vụ", "Mô tả"],
        [
            ["Mua hàng", "Khách xem sản phẩm, thêm giỏ, áp voucher, checkout, thanh toán VNPay và theo dõi đơn."],
            ["Thuê thiết bị", "Khách kiểm tra khả dụng, tạo đơn thuê, ký hợp đồng, thanh toán phí thuê; Staff xử lý bàn giao, cọc, trả thiết bị và hoàn tất."],
            ["Xác thực eKYC", "Khách upload giấy tờ/selfie, gửi hồ sơ; Admin/Staff duyệt hoặc từ chối."],
            ["Quản trị vận hành", "Quản lý catalog, tồn kho, đơn hàng, voucher, review, support ticket, audit log."],
            ["Phân quyền", "Super Admin quản lý user, role, permission và gán quyền."],
        ],
    )
    add_para(
        doc,
        "Các rủi ro nghiệp vụ chính gồm rò rỉ dữ liệu định danh, truy cập trái phép vào đơn hàng/eKYC, sai lệch trạng thái thanh toán, gian lận thuê thiết bị và thiếu khả năng truy vết thao tác nội bộ. Do đó hệ thống cần kiến trúc bảo mật nhiều lớp gồm xác thực, phân quyền, kiểm soát upload, xác thực callback thanh toán và audit log.",
    )

    add_heading(doc, "2.2. PHÂN TÍCH YÊU CẦU HỆ THỐNG", 2)
    add_heading(doc, "2.2.1. Yêu cầu chức năng", 3)
    add_table(
        doc,
        ["ID", "Vai trò", "Tên chức năng", "Mô tả"],
        [
            ["G-01", "Guest", "Xem catalog", "Xem trang chủ, danh sách sản phẩm, chi tiết sản phẩm, review và tình trạng có thể thuê."],
            ["G-02", "Guest", "Đăng ký/kích hoạt", "Tạo tài khoản và kích hoạt qua email."],
            ["G-03", "Guest", "Gửi hỗ trợ", "Gửi support ticket công khai."],
            ["C-01", "Customer", "Xác thực", "Đăng nhập, refresh session, đăng xuất, quên/reset mật khẩu."],
            ["C-02", "Customer", "Quản lý hồ sơ", "Cập nhật thông tin cá nhân, avatar và địa chỉ."],
            ["C-03", "Customer", "Mua hàng", "Quản lý giỏ, áp voucher, checkout, thanh toán VNPay, theo dõi đơn."],
            ["C-04", "Customer", "Thuê thiết bị", "Tạo đơn thuê, ký hợp đồng, thanh toán phí thuê và theo dõi đơn thuê."],
            ["C-05", "Customer", "eKYC", "Upload giấy tờ/selfie, submit hồ sơ và xem trạng thái duyệt."],
            ["A-01", "Staff/Admin", "Quản lý catalog", "Tạo/sửa/xóa/khôi phục sản phẩm, danh mục, ảnh, gallery, giá."],
            ["A-02", "Staff/Admin", "Quản lý tồn kho", "Điều chỉnh tồn kho bán/thuê/tổng và xem audit log tồn kho."],
            ["A-03", "Staff/Admin", "Quản lý đơn", "Xem/cập nhật đơn mua, xử lý vòng đời đơn thuê."],
            ["A-04", "Staff/Admin", "Quản trị hỗ trợ", "Xử lý support ticket, review, eKYC và audit log."],
            ["SA-01", "Super Admin", "Quản lý RBAC", "Quản lý user, role, permission và gán permission cho role."],
        ],
    )
    add_heading(doc, "2.2.2. Yêu cầu phi chức năng", 3)
    add_table(
        doc,
        ["Nhóm yêu cầu", "Tiêu chí"],
        [
            ["Hiệu năng", "Các API danh sách hỗ trợ phân trang/lọc; thao tác thường phản hồi nhanh trong điều kiện tải thông thường."],
            ["Bảo mật", "JWT, refresh token, BCrypt, phân quyền bằng permission, xác thực callback VNPay, không ghi log secret."],
            ["Bảo vệ dữ liệu", "Kiểm soát quyền xem dữ liệu cá nhân/eKYC; kiểm soát upload file; audit thao tác nhạy cảm."],
            ["Khả dụng", "Giao diện rõ ràng cho Customer và Admin; form có validate, loading, lỗi và trạng thái nghiệp vụ."],
            ["Mở rộng", "Backend tách module theo domain; frontend tách service theo tài nguyên; dễ bổ sung module mới."],
            ["Bảo trì", "Dùng DTO, service layer, repository layer; tài liệu cập nhật theo thay đổi nghiệp vụ."],
        ],
    )

    add_heading(doc, "2.3. MÔ HÌNH HÓA YÊU CẦU", 2)
    add_table(
        doc,
        ["Actor", "Mô tả quyền hạn"],
        [
            ["Guest", "Xem sản phẩm, review, kiểm tra khả dụng thuê, đăng ký và gửi hỗ trợ."],
            ["Customer", "Mua/thuê thiết bị, thanh toán, quản lý hồ sơ, địa chỉ, eKYC, review."],
            ["Staff", "Vận hành đơn hàng, đơn thuê, catalog, tồn kho, support và eKYC theo quyền được cấp."],
            ["Admin", "Quản lý vận hành hệ thống ở phạm vi rộng hơn Staff."],
            ["Super Admin", "Quản lý user, role, permission và phân quyền."],
            ["VNPay", "Xử lý thanh toán và trả kết quả thanh toán."],
            ["Mail Service", "Gửi email kích hoạt tài khoản và reset mật khẩu."],
            ["Redis", "Hỗ trợ blacklist token/cache phiên làm việc."],
        ],
    )
    fig = add_figure(doc, diagram["02-usecase-overview"][0], diagram["02-usecase-overview"][1], fig)
    add_heading(doc, "Đặc tả các Use Case quan trọng", 3)
    add_table(
        doc,
        ["Use Case", "Pre-condition", "Basic Flow", "Alternate Flow", "Post-condition"],
        [
            ["Đăng nhập", "User có tài khoản hợp lệ.", "Nhập email/password; backend xác thực; sinh token; frontend điều hướng theo role.", "Sai mật khẩu, chưa kích hoạt, bị khóa, token hết hạn.", "User đăng nhập và có quyền tương ứng."],
            ["Checkout mua hàng", "Customer đăng nhập, có giỏ hàng và địa chỉ.", "Chọn item, áp voucher, tạo order, tạo VNPay URL, thanh toán, cập nhật trạng thái.", "Voucher/tồn kho/payment lỗi.", "Đơn mua và payment log được cập nhật."],
            ["Thuê thiết bị", "Customer đăng nhập, thiết bị khả dụng.", "Check lịch, checkout, ký hợp đồng, thanh toán, Staff bàn giao, trả thiết bị, hoàn tất.", "Thiết bị không khả dụng, thanh toán lỗi, phát sinh hư hỏng.", "Đơn thuê hoàn tất hoặc ở trạng thái phù hợp."],
            ["eKYC", "Customer đăng nhập.", "Initiate, upload front/back/selfie, submit, Admin duyệt.", "Thiếu file, file lỗi, hồ sơ bị từ chối.", "Trạng thái xác minh được cập nhật."],
            ["Quản lý RBAC", "Super Admin đăng nhập.", "Tạo role/permission, gán permission, cập nhật user.", "Trùng tên role/permission, thiếu quyền.", "Quyền truy cập hệ thống được cập nhật."],
        ],
    )

    add_heading(doc, "2.4. PHÂN TÍCH MỐI ĐE DỌA HỆ THỐNG", 2)
    add_table(
        doc,
        ["Mối đe dọa", "Kịch bản", "Hậu quả", "Biện pháp"],
        [
            ["Rò rỉ eKYC", "Người không có quyền truy cập ảnh giấy tờ/selfie.", "Lộ dữ liệu cá nhân nghiêm trọng.", "RBAC, kiểm soát file access, audit log."],
            ["Token theft", "Access token bị đánh cắp.", "Truy cập trái phép.", "JWT ngắn hạn, refresh token, blacklist Redis."],
            ["Broken access control", "User thường gọi API quản trị.", "Lộ hoặc sửa dữ liệu nhạy cảm.", "@PreAuthorize và permission model."],
            ["Payment spoofing", "Giả callback VNPay.", "Sai trạng thái thanh toán.", "Xác minh chữ ký/hash và lưu payment log."],
            ["File upload attack", "Upload file độc hại hoặc path traversal.", "Rủi ro thực thi/lộ file.", "Validate type/size/path, isolate upload."],
            ["Insider misuse", "Nhân sự nội bộ xem/sửa dữ liệu ngoài phạm vi.", "Rò rỉ hoặc sai lệch dữ liệu.", "Least privilege và audit log."],
        ],
    )

    add_heading(doc, "2.5. THIẾT KẾ KIẾN TRÚC HỆ THỐNG", 2)
    fig = add_figure(doc, diagram["01-architecture"][0], diagram["01-architecture"][1], fig)
    add_table(
        doc,
        ["Thành phần", "Công nghệ", "Vai trò"],
        [
            ["Frontend", "Next.js, React, TypeScript, Tailwind/shadcn", "Giao diện Customer/Admin, gọi API."],
            ["Backend", "Java 17, Spring Boot", "REST API, nghiệp vụ, bảo mật, tích hợp ngoài."],
            ["Database", "PostgreSQL", "Lưu user, sản phẩm, đơn, thuê, eKYC, voucher, audit."],
            ["Redis", "Redis", "Token blacklist/cache."],
            ["Payment", "VNPay", "Thanh toán đơn mua và đơn thuê."],
            ["Mail", "SMTP", "Kích hoạt tài khoản, reset mật khẩu."],
        ],
    )

    add_heading(doc, "2.6. THIẾT KẾ CHI TIẾT", 2)
    for slug in [
        "03-activity-auth",
        "04-activity-checkout",
        "05-activity-rental",
        "06-sequence-auth",
        "07-sequence-vnpay",
        "08-sequence-ekyc",
        "09-sequence-rental",
        "10-sequence-admin-catalog-inventory",
        "11-sequence-rbac",
    ]:
        fig = add_figure(doc, diagram[slug][0], diagram[slug][1], fig)
    add_para(
        doc,
        "Các sơ đồ hoạt động và tuần tự tập trung vào những luồng nghiệp vụ có nhiều tương tác giữa người dùng, Frontend, Backend, Database và dịch vụ ngoài. Những luồng đơn giản như dashboard widget, reset mật khẩu chi tiết hoặc support ticket được mô tả ở mức use case để tránh làm tài liệu quá rời rạc.",
    )

    add_heading(doc, "2.7. THIẾT KẾ BẢO MẬT DỮ LIỆU", 2)
    add_para(
        doc,
        "Trong phạm vi chương này, hệ thống tập trung vào các cơ chế bảo mật dữ liệu đã thể hiện trong kiến trúc hiện tại: xác thực JWT, refresh token, phân quyền theo permission, mã hóa mật khẩu bằng BCrypt, blacklist token bằng Redis, kiểm soát quyền truy cập API bằng @PreAuthorize, kiểm tra callback thanh toán và audit log. Phần mã hóa E2EE chi tiết không được trình bày trong Chương 2 theo phạm vi đã chốt.",
    )
    add_bullets(
        doc,
        [
            "Dữ liệu tài khoản được bảo vệ qua mật khẩu đã băm và cơ chế token.",
            "Dữ liệu định danh/eKYC chỉ được truy cập bởi người dùng hợp lệ hoặc nhân sự có quyền xét duyệt.",
            "Các thao tác quản trị quan trọng cần được ghi nhận trong audit log để phục vụ truy vết.",
            "Thông tin thanh toán được xử lý qua VNPay và xác minh bằng callback hợp lệ.",
        ],
    )

    add_heading(doc, "2.8. THIẾT KẾ CƠ SỞ DỮ LIỆU", 2)
    fig = add_figure(doc, diagram["12-erd-logical"][0], diagram["12-erd-logical"][1], fig)
    add_table(
        doc,
        ["Nhóm dữ liệu", "Thực thể chính"],
        [
            ["Auth/RBAC", "User, UserProfile, Role, Permission, RefreshToken, ActivationToken, PasswordResetToken"],
            ["Catalog", "Product, ProductImage, ProductSpecification, ProductPriceHistory, Category"],
            ["Cart/Order", "CartItem, Order, OrderItem"],
            ["Rental", "RentalOrder, RentalOrderItem, RentalContract, Device, HandoverReport, ReturnReport, RentalPayment"],
            ["Payment/Voucher", "PaymentTransactionLog, Voucher"],
            ["eKYC", "UserIdentity, VerificationSession, VerificationArtifact, VerificationResult, RiskAssessment"],
            ["Support/Audit", "SupportTicket, AuditLog, InventoryAuditLog"],
        ],
    )
    add_table(
        doc,
        ["Bảng", "Cột tiêu biểu", "Ý nghĩa"],
        [
            ["users", "id, email, password, status", "Thông tin tài khoản và trạng thái người dùng."],
            ["products", "id, name, price, stock, category_id", "Thông tin thiết bị máy ảnh/ống kính."],
            ["orders", "id, user_id, status, payment_status", "Đơn mua hàng."],
            ["rental_orders", "id, user_id, status, rental_period", "Đơn thuê thiết bị."],
            ["devices", "id, product_id, serial, status", "Thiết bị vật lý cho thuê."],
            ["verification_sessions", "id, user_id, status", "Phiên xác thực eKYC."],
            ["vouchers", "id, code, discount, status", "Mã khuyến mãi."],
            ["audit_logs", "id, actor_id, action, created_at", "Nhật ký thao tác hệ thống."],
        ],
    )

    add_heading(doc, "2.9. THIẾT KẾ GIAO DIỆN HỆ THỐNG", 2)
    fig = add_figure(doc, diagram["13-sitemap"][0], diagram["13-sitemap"][1], fig)
    add_table(
        doc,
        ["Màn hình", "Thành phần chính", "Mục đích"],
        [
            ["Trang chủ/catalog", "Danh sách sản phẩm, filter, CTA xem chi tiết", "Giúp Guest/Customer tìm thiết bị."],
            ["Chi tiết sản phẩm", "Ảnh, thông số, giá, review, add cart/rent", "Cung cấp thông tin trước khi mua/thuê."],
            ["Cart/Checkout", "Item, địa chỉ, voucher, tổng tiền, nút thanh toán", "Hoàn tất đơn mua."],
            ["Profile/eKYC", "Thông tin cá nhân, upload giấy tờ, trạng thái duyệt", "Quản lý tài khoản và định danh."],
            ["Admin dashboard", "Revenue, order stats, top products, low stock", "Theo dõi vận hành."],
            ["Admin product/inventory", "Table, search, form dialog, stock action", "Quản lý catalog và tồn kho."],
            ["Admin rental", "Danh sách đơn thuê, trạng thái, action bàn giao/trả", "Vận hành vòng đời thuê."],
            ["Super Admin RBAC", "Role, permission, user assignment", "Quản lý phân quyền."],
        ],
    )


def replace_chapter2(diagrams: list[tuple[str, str, Path]]) -> None:
    doc = Document(str(SOURCE_DOCX))
    for style_name in ["Normal"]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(13)

    paragraphs = doc.paragraphs
    start_p = None
    end_p = None
    pairs = []
    current_start = None
    for p in paragraphs:
        text = p.text.strip()
        if text.startswith("Chương 2: PHÂN TÍCH"):
            current_start = p
            continue
        if current_start is not None and text.startswith("Chương 3:"):
            pairs.append((current_start, p))
            current_start = None
    if pairs:
        start_p, end_p = pairs[-1]
    if start_p is None or end_p is None:
        raise RuntimeError("Cannot locate Chapter 2 / Chapter 3 boundaries in template.")

    body = doc.element.body
    children = list(body)
    start_idx = children.index(start_p._p)
    end_idx = children.index(end_p._p)
    for el in children[start_idx:end_idx]:
        body.remove(el)

    anchor = end_p._p
    before_ids = {id(el) for el in body}
    build_chapter2(doc, diagrams)
    new_elements = [el for el in body if id(el) not in before_ids]
    for el in new_elements:
        anchor.addprevious(el)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_DOCX))


def main() -> None:
    ensure_dirs()
    diagrams = write_diagrams()
    exported = export_diagrams(diagrams)
    replace_chapter2(exported)
    print(f"DOCX={OUT_DOCX}")
    print(f"DIAGRAM_DIR={DIAGRAM_DIR}")


if __name__ == "__main__":
    main()
